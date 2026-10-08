from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path
import csv
import json
import os
import tempfile
import unittest

from openpyxl import load_workbook
from finance.logic import ReportError, analyze_sources, same_employer, placement_status
from finance.service import ROOT, generate, settings
from finance.uploads import read_source, common_export_date, scalar
from finance.exporter import Template


def examples():
    return {kind: {"data": (ROOT / "examples" / name).read_bytes(), "name": name, "sheet": None, "header_row": 1} for kind, name in (("students", "Students_2026_09_30.csv"), ("internships", "Internship_Placements_2026_09_30.csv"), ("offers", "Full_Time_Placements_2026_09_30.csv"))}


def source_rows():
    return {kind: read_source(item["data"], item["name"], kind) for kind, item in examples().items()}


class Rules(unittest.TestCase):
    def setUp(self):
        self.scope, _ = settings()
        self.sources = source_rows()

    def test_exact_populations(self):
        model = analyze_sources(self.sources, self.scope, "2026-09-30")
        self.assertEqual(len(model["trackerRows"]), 4)
        self.assertEqual(len(model["historyRows"]), 5)
        self.assertEqual(sum(s["placed"] for s in model["trackerStats"]), 1)
        self.assertEqual(model["futureAdmitCount"], 1)
        self.assertTrue(all(row[6] in self.scope["trackerTerms"] for row in model["trackerRows"]))
        self.assertTrue(all(row[0] in (2024, 2025, 2026) for row in model["historyRows"]))

    def test_completed_actual_graduates_are_excluded(self):
        row = next(s for s in self.sources["students"] if s["netID"] == "demo_seek")
        row["Actual Graduation Semester"] = "Summer 2026"
        model = analyze_sources(self.sources, self.scope, "2026-09-30")
        self.assertNotIn("demo_seek", [r[0] for r in model["trackerRows"]])

    def test_status_mapping_and_unknown_status(self):
        self.assertEqual(placement_status("Not seeking: Further education"), "Not Seeking")
        self.assertEqual(placement_status("Not Reported"), "no status")
        with self.assertRaises(ReportError):
            placement_status("Unknown new status")

    def test_company_ids_take_priority(self):
        i = {"Company.id": "1", "Company": "Same Name"}
        o = {"Company.id": "2", "Company": "Same Name"}
        self.assertFalse(same_employer(i, o))
        o["Company.id"] = ""
        o["Company"] = "same-name"
        self.assertTrue(same_employer(i, o))

    def test_future_acceptance_excluded(self):
        offer = next(o for o in self.sources["offers"] if o["Accepted?"].lower() == "true")
        offer["Accepted Date"] = "2026-10-30"
        model = analyze_sources(self.sources, self.scope, "2026-09-30")
        for detail in model["trackerDetails"] + model["historyDetails"]:
            self.assertNotIn(offer["Record Id"], [o["Record Id"] for o in detail["accepted"]])

    def test_duplicate_record_ids_rejected(self):
        self.sources["students"].append(deepcopy(self.sources["students"][0]))
        with self.assertRaisesRegex(ReportError, "duplicate Record Id"):
            analyze_sources(self.sources, self.scope, "2026-09-30")

    def test_unmatched_student_ids_rejected(self):
        self.sources["internships"][0]["Student ID.id"] = "no-match"
        with self.assertRaisesRegex(ReportError, "no matching"):
            analyze_sources(self.sources, self.scope, "2026-09-30")

    def test_invalid_dates_rejected(self):
        self.sources["internships"][0]["Start Date"] = "2026-02-30"
        with self.assertRaises(ReportError):
            analyze_sources(self.sources, self.scope, "2026-09-30")

    def test_long_excel_numeric_identifiers_rejected(self):
        with self.assertRaisesRegex(ReportError, "as text"):
            scalar(123456789012345678, "Record Id")
        self.assertEqual(scalar(123, "Phone", "000000"), "000123")

    def test_mixed_export_dates_rejected(self):
        with self.assertRaisesRegex(ReportError, "different export dates"):
            common_export_date(["Students_2026_09_30.xlsx", "Internship_Placements_2026_10_01.xlsx"])


class Generation(unittest.TestCase):
    def test_layout_cached_rates_and_exclusions(self):
        result = generate(examples(), "2026-09-30")
        files = [data for name, data in result["files"].items() if name.endswith(".xlsx")]
        tracker = load_workbook(BytesIO(files[0]), data_only=True)
        history = load_workbook(BytesIO(files[1]), data_only=True)
        self.assertEqual(tracker.sheetnames, ["Placement Summary", "Placement Tracker", "Review"])
        self.assertEqual(tracker["Placement Summary"]["B10"].value, 4)
        self.assertEqual(tracker["Placement Summary"]["D10"].value, 0.25)
        self.assertEqual(tracker["Placement Tracker"].freeze_panes, "D9")
        self.assertEqual(history["Placement Summary"]["D8"].value, "—")
        self.assertEqual(history["Student Detail"].freeze_panes, "E7")
        tracker.close(); history.close()

    def test_combined_excel_matches_csv(self):
        data = (ROOT / "examples/Finance_Source_Exports_2026_09_30.xlsx").read_bytes()
        inputs = {kind: {"data": data, "name": "Finance_Source_Exports_2026_09_30.xlsx", "sheet": sheet, "header_row": 1} for kind, sheet in (("students", "Students"), ("internships", "Internship Placements"), ("offers", "Full-Time Placements"))}
        excel = generate(inputs, "2026-09-30")
        csv_result = generate(examples(), "2026-09-30")
        for key in ("trackerRows", "historyRows", "reviewRows", "trackerStats", "historicalStats"):
            self.assertEqual(excel["model"][key], csv_result["model"][key])

    def test_previous_tracker_has_no_changes(self):
        inputs = examples()
        first = generate(inputs, "2026-09-30")
        name = next(name for name in first["files"] if "Tracker" in name)
        second = generate(inputs, "2026-09-30", {"data": first["files"][name], "name": name})
        self.assertEqual(second["changes"]["rows"], [])

    def test_empty_exports_have_dash_rates(self):
        inputs = examples()
        for item in inputs.values():
            item["data"] = (item["data"].decode("utf-8-sig").splitlines()[0] + "\n").encode()
        result = generate(inputs, "2026-09-30")
        self.assertEqual(result["manifest"]["counts"]["tracker"], 0)
        data = next(data for name, data in result["files"].items() if "Tracker" in name)
        book = load_workbook(BytesIO(data), data_only=True)
        self.assertEqual(book["Placement Summary"]["D10"].value, "—")
        book.close()

    def test_missing_columns_and_misdated_reports_rejected(self):
        inputs = examples()
        inputs["students"]["data"] = b"First Name,Last Name\nDemo,Example\n"
        with self.assertRaisesRegex(ReportError, "required columns"):
            generate(inputs, "2026-09-30")
        with self.assertRaisesRegex(ReportError, "match the export date"):
            generate(examples(), "2026-10-08")

    def test_format_tamper_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            lock = json.loads((ROOT / "format-lock.json").read_text())
            for name in lock["sha256"]:
                (root / name).parent.mkdir(parents=True, exist_ok=True)
                (root / name).write_bytes((ROOT / name).read_bytes())
            (root / "format-lock.json").write_text(json.dumps(lock))
            (root / "format.json").write_text("{}")
            with self.assertRaisesRegex(ReportError, "format check"):
                settings(root)

    def test_source_formula_text_stays_literal(self):
        inputs = examples()
        text = inputs["students"]["data"].decode()
        text = text.replace("demo_seek@example.test", "=1+1")
        inputs["students"]["data"] = text.encode()
        result = generate(inputs, "2026-09-30")
        data = next(data for name, data in result["files"].items() if "Tracker" in name)
        book = load_workbook(BytesIO(data), data_only=False)
        cell = next(row[7] for row in book["Placement Tracker"].iter_rows(min_row=9, max_col=10) if row[0].value == "demo_seek")
        self.assertEqual(cell.value, "=1+1")
        self.assertNotEqual(cell.data_type, "f")
        book.close()

    def test_review_status_and_date_are_preserved_only_for_unchanged_issue(self):
        template = Template(ROOT / "assets/tracker_template.xlsx")
        template.set(2, "G7", "Verified")
        template.set(2, "H7", (date(2026, 9, 30) - date(1899, 12, 30)).days)
        previous = {"data": template.finish(), "name": "Previous_Tracker.xlsx"}
        result = generate(examples(), "2026-09-30", previous)
        first_review = result["model"]["reviewRows"][0]
        self.assertEqual(first_review[-2:], ["Verified", "2026-09-30"])
        template.set(2, "E7", "A different issue")
        previous["data"] = template.finish()
        result = generate(examples(), "2026-09-30", previous)
        self.assertEqual(result["model"]["reviewRows"][0][-2:], ["Pending", None])

    def test_custom_csv_header_row(self):
        item = examples()["students"]
        data = b"Source export\nPrepared for Finance\n" + item["data"]
        rows = read_source(data, item["name"], "students", header_row=3)
        self.assertEqual(len(rows), len(source_rows()["students"]))


class StreamlitFlow(unittest.TestCase):
    def app(self):
        from streamlit.testing.v1 import AppTest
        self.assertFalse(os.environ.get("APP_PASSWORD"), "Run normal UI tests without APP_PASSWORD.")
        return AppTest.from_file(str(ROOT / "app.py"), default_timeout=20).run()

    def test_empty_app_and_demo(self):
        app = self.app()
        self.assertEqual(len(app.exception), 0)
        self.assertTrue(app.button(key="generate").disabled)
        app.radio(key="mode").set_value("Try a fictional example").run()
        app.button(key="generate").click().run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.error), 0)
        self.assertEqual(app.metric[0].value, "4")

    def test_three_csv_uploads(self):
        app = self.app()
        for kind, item in examples().items():
            app.file_uploader(key=f"upload_{kind}_0").upload(item["name"], item["data"], "text/csv").run()
        self.assertFalse(app.button(key="generate").disabled)
        app.button(key="generate").click().run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.error), 0)
        self.assertEqual(app.metric[2].value, "25.0%")
        # Any input change must hide previous downloads until regeneration.
        app.file_uploader(key="upload_students_0").clear().run()
        self.assertEqual(len(app.metric), 0)
        self.assertTrue(app.button(key="generate").disabled)

    def test_excel_workbook_upload(self):
        app = self.app()
        app.radio(key="mode").set_value("One Excel workbook").run()
        name = "Finance_Source_Exports_2026_09_30.xlsx"
        app.file_uploader(key="combined_0").upload(name, (ROOT / "examples" / name).read_bytes()).run()
        self.assertFalse(app.button(key="generate").disabled)
        app.button(key="generate").click().run()
        self.assertEqual(len(app.error), 0)
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.metric[0].value, "4")

    def test_password_gate(self):
        from unittest.mock import patch
        from streamlit.testing.v1 import AppTest
        with patch.dict(os.environ, {"APP_PASSWORD": "unit-test-password"}):
            app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=20).run()
            self.assertEqual(len(app.file_uploader), 0)
            app.text_input[0].set_value("wrong")
            app.button[0].click().run()
            self.assertEqual(len(app.file_uploader), 0)
            self.assertEqual(len(app.error), 1)
            app.text_input[0].set_value("unit-test-password")
            app.button[0].click().run()
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.file_uploader), 4)


if __name__ == "__main__":
    unittest.main()
