"""Read-only Excel/CSV uploads. Uploaded bytes never need a server-side file."""
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
import csv
import re
import zipfile

from openpyxl import load_workbook

from .logic import ReportError, REQUIRED_COLUMNS, validate_columns, valid_date

MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 200
ID_FIELDS = {"Record Id", "Student ID.id", "StudentID.id", "Company.id", "netID"}
DATE_FIELDS = {"Start Date", "Offer Date", "Accepted Date"}


def check_file(data, filename):
    if len(data) > MAX_BYTES:
        raise ReportError(f"{Path(filename).name}: each upload must be 25 MB or smaller.")
    extension = Path(filename).suffix.lower()
    if extension not in (".xlsx", ".csv"):
        raise ReportError("Upload an .xlsx or .csv file. Save older .xls workbooks as .xlsx first.")
    if extension == ".xlsx":
        try:
            with zipfile.ZipFile(BytesIO(data)) as book:
                if sum(i.file_size for i in book.infolist()) > 250 * 1024 * 1024:
                    raise ReportError("The expanded Excel workbook is too large.")
                if "xl/workbook.xml" not in book.namelist():
                    raise ReportError("This file is not an Excel workbook.")
        except zipfile.BadZipFile as error:
            raise ReportError("The Excel file is damaged or password protected. Upload an unencrypted .xlsx export.") from error
    return extension


def sheet_names(data, filename):
    if check_file(data, filename) == ".csv":
        return []
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=False)
        names = book.sheetnames
        book.close()
        return names
    except (ValueError, OSError, KeyError, TypeError) as error:
        raise ReportError("The Excel file could not be opened. Save an unencrypted .xlsx copy and upload it.") from error


def filename_date(filename):
    match = re.search(r"(?<!\d)(\d{4})[_-](\d{2})[_-](\d{2})(?!\d)", Path(filename).name)
    value = "-".join(match.groups()) if match else None
    if value and not valid_date(value):
        raise ReportError(f"{Path(filename).name}: invalid date in the filename.")
    return date.fromisoformat(value) if value else None


def common_export_date(filenames):
    dates = {filename_date(name) for name in filenames} - {None}
    if len(dates) > 1:
        raise ReportError("The source filenames have different export dates. Upload a matching-date set.")
    return next(iter(dates)) if dates else None


def scalar(value, field="", number_format="General"):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        if field in ID_FIELDS and abs(value) >= 10**15:
            raise ReportError(f"{field}: a long identifier is stored as a number. Export identifiers as text to preserve every digit.")
        if isinstance(value, float) and not value.is_integer():
            return str(value)
        value = int(value)
        if re.fullmatch(r"0+", number_format or ""):
            return str(value).zfill(len(number_format))
        return str(value)
    return str(value).strip()


def source_date(value, field, row_number):
    if not value:
        return ""
    if valid_date(value):
        return value
    for fmt in ("%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    if "T" in value:
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).date().isoformat()
        except ValueError:
            pass
    raise ReportError(f"{field}: invalid date on source row {row_number}. Use real Excel dates, YYYY-MM-DD, or MM/DD/YYYY.")


def _csv_rows(data):
    try:
        text = data.decode("utf-16") if data.startswith((b"\xff\xfe", b"\xfe\xff")) else data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("cp1252")
        except UnicodeDecodeError as error:
            raise ReportError("Save this CSV as UTF-8 and upload it again.") from error
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    return csv.reader(StringIO(text), dialect)


def read_source(data, filename, kind, sheet=None, header_row=1):
    extension = check_file(data, filename)
    book = None
    try:
        if not 1 <= header_row <= 25:
            raise ReportError("The header row must be between 1 and 25.")
        if extension == ".xlsx":
            book = load_workbook(BytesIO(data), read_only=True, data_only=False)
            if sheet not in book.sheetnames:
                raise ReportError(f"Choose the {kind} worksheet.")
            ws = book[sheet]
            if ws.max_row and ws.max_row > MAX_ROWS + header_row:
                raise ReportError(f"{kind}: more than {MAX_ROWS:,} rows. Export a smaller source table.")
            iterator = ws.iter_rows(max_col=min(ws.max_column or MAX_COLUMNS, MAX_COLUMNS))
        else:
            iterator = iter(_csv_rows(data))
        for _ in range(header_row - 1):
            next(iterator, None)
        heading_cells = next(iterator, None)
        if heading_cells is None:
            raise ReportError(f"{kind}: the header row is missing.")
        headers = [scalar(c.value if extension == ".xlsx" else c).lstrip("\ufeff") for c in heading_cells]
        while headers and not headers[-1]:
            headers.pop()
        validate_columns(headers, kind)
        fields = set(REQUIRED_COLUMNS[kind]) | ({"Student ID.id", "StudentID.id"} if kind == "offers" else set())
        result = []
        for row_number, cells in enumerate(iterator, start=header_row + 1):
            if row_number - header_row > MAX_ROWS:
                raise ReportError(f"{kind}: more than {MAX_ROWS:,} rows.")
            if not any(c.value is not None if extension == ".xlsx" else bool(c) for c in cells):
                continue
            if len(cells) > len(headers) and any(c.value is not None if extension == ".xlsx" else bool(c) for c in cells[len(headers):]):
                raise ReportError(f"{kind}: values without a column heading on row {row_number}.")
            record = {name: "" for name in headers}
            for position, header in enumerate(headers):
                if position >= len(cells):
                    break
                cell = cells[position]
                if extension == ".xlsx":
                    if cell.data_type == "f" and header in fields:
                        raise ReportError(f"{kind}: {header} contains a formula on row {row_number}. Upload source values, or paste values in an export copy.")
                    record[header] = scalar(cell.value, header, cell.number_format)
                else:
                    record[header] = scalar(cell, header)
                if header in DATE_FIELDS:
                    record[header] = source_date(record[header], header, row_number)
            record["sourceRow"] = row_number
            result.append(record)
        return result
    except ReportError:
        raise
    except (ValueError, OSError, KeyError, TypeError, csv.Error) as error:
        raise ReportError(f"{kind}: could not read this source table. Check the worksheet, header row, and file format.") from error
    finally:
        if book is not None:
            book.close()


def read_previous(data, filename, format_settings):
    if check_file(data, filename) != ".xlsx":
        raise ReportError("The previous tracker must be an .xlsx report generated in this format.")
    book = load_workbook(BytesIO(data), read_only=True, data_only=False)
    try:
        if not {"Placement Tracker", "Review"}.issubset(book.sheetnames):
            raise ReportError("The previous tracker must contain Placement Tracker and Review tabs.")
        tracker = book["Placement Tracker"]
        headers = [c.value for c in next(tracker.iter_rows(min_row=8, max_row=8, max_col=10))]
        if headers != format_settings["headers"]["tracker"]:
            raise ReportError("The previous tracker has different column headings. Use a report from the approved format.")
        rows = [[scalar(c.value) for c in cells] for cells in tracker.iter_rows(min_row=9, max_col=10) if cells[0].value]
        if len({r[0].lower() for r in rows}) != len(rows):
            raise ReportError("The previous tracker contains duplicate NetIDs.")
        review = book["Review"]
        review_headers = [c.value for c in next(review.iter_rows(min_row=6, max_row=6, max_col=8))]
        if review_headers != format_settings["headers"]["review"]:
            raise ReportError("The previous Review tab has different column headings.")
        prior_review = {}
        for cells in review.iter_rows(min_row=7, max_col=8):
            if not cells[0].value:
                continue
            row = [scalar(c.value) for c in cells]
            if row[0] in prior_review:
                raise ReportError("The previous Review tab contains duplicate NetIDs.")
            if row[6] not in ("Pending", "Verified", "Corrected"):
                raise ReportError("Previous Review status must be Pending, Verified, or Corrected.")
            review_date = row[7]
            if review_date:
                review_date = source_date(review_date, "Reviewed on", cells[7].row)
            prior_review[row[0]] = {"issue": row[4], "status": row[6], "date": review_date or None}
        return rows, prior_review
    finally:
        book.close()
